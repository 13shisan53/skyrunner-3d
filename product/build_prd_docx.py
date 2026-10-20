"""Convert the SKYRUNNER PRD (Markdown, known subset) into a styled DOCX."""

import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_DEFAULT_DIR = r"C:\Users\Gu mo\Documents\deepseek-harness\default-workspace\skyrunner-3d\product"
SRC = sys.argv[1] if len(sys.argv) > 1 else _DEFAULT_DIR + r"\PRD.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else _DEFAULT_DIR + r"\SKYRUNNER-PRD.docx"
TITLE_ARG = sys.argv[3] if len(sys.argv) > 3 else None

INK = RGBColor(0x1B, 0x22, 0x2B)
MUTED = RGBColor(0x5A, 0x67, 0x75)
ACCENT = RGBColor(0x0E, 0x7C, 0x86)
ACCENT_HEX = "0E7C86"
HEADER_FILL = "E8F4F5"
CODE_FILL = "F3F5F7"
RULE_HEX = "D6DEE5"

CN_FONT = "微软雅黑"
MONO_FONT = "Consolas"
CONTENT_W_CM = 16.4


def set_run_font(run, name, size, bold=False, color=INK, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def shade(element_pr, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    element_pr.append(shd)


def para_shade(paragraph, hex_fill):
    shade(paragraph._p.get_or_add_pPr(), hex_fill)


def para_border(paragraph, side="bottom", hex_color=RULE_HEX, size=6, space=6):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), hex_color)
    borders.append(border)
    ppr.append(borders)


def cell_shade(cell, hex_fill):
    shade(cell._tc.get_or_add_tcPr(), hex_fill)


def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def table_borders(table, hex_color="C8D3DC"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), hex_color)
        borders.append(el)
    tbl_pr.append(borders)


def no_split_rows(table):
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_field(paragraph, instr, font_size=8.5, color=MUTED):
    run = paragraph.add_run()
    set_run_font(run, CN_FONT, font_size, color=color)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_end)


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\n]+\*)")
URL_RE = re.compile(r"(https?://[^\s\u4e00-\u9fff，。、；：）】]+)")


def add_hyperlink(paragraph, url, size, bold=False):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)

    run_el = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), CN_FONT)
    rfonts.set(qn("w:hAnsi"), CN_FONT)
    rfonts.set(qn("w:eastAsia"), CN_FONT)
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0E7C86")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    if bold:
        rpr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)

    run_el.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = url
    run_el.append(t)
    hl.append(run_el)
    paragraph._p.append(hl)


def add_inline(paragraph, text, size=10.5, color=INK, bold=False):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, CN_FONT, size, bold=True, color=color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, CN_FONT, size, bold=bold, color=color, italic=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, MONO_FONT, size - 1, color=RGBColor(0x0B, 0x5C, 0x66))
            run.font.bold = bold
        else:
            for chunk in URL_RE.split(part):
                if not chunk:
                    continue
                if URL_RE.fullmatch(chunk):
                    add_hyperlink(paragraph, chunk, size, bold)
                else:
                    run = paragraph.add_run(chunk)
                    set_run_font(run, CN_FONT, size, bold=bold, color=color)


def parse_blocks(md):
    lines = md.split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if not s:
            i += 1
            continue

        if s.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(code)))
            continue

        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            blocks.append(("table", rows))
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            blocks.append((f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue

        if s.startswith("> "):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(("quote", " ".join(buf)))
            continue

        if re.match(r"^-\s+", s):
            blocks.append(("bullet", re.sub(r"^-\s+", "", s)))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            blocks.append(("num", f"{m.group(1)}. {m.group(2)}"))
            i += 1
            continue

        if re.match(r"^-{3,}$", s):
            blocks.append(("hr", ""))
            i += 1
            continue

        blocks.append(("p", s))
        i += 1
    return blocks


def split_table_row(row):
    inner = row.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [c.strip() for c in inner.split("|")]


def is_separator(cells):
    return all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells if c != "")


def build_table(doc, rows):
    header = split_table_row(rows[0])
    body = []
    for r in rows[1:]:
        cells = split_table_row(r)
        if is_separator(cells):
            continue
        body.append(cells)

    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table)
    cell_margins(table)

    def visual_len(text):
        n = 0
        for ch in text:
            n += 2 if "\u4e00" <= ch <= "\u9fff" else 1
        return n

    weights = []
    for c in range(ncols):
        longest = visual_len(header[c])
        for r in body:
            if c < len(r):
                longest = max(longest, visual_len(r[c]))
        weights.append(min(max(longest, 4), 70))

    # every column gets a floor width, the rest is shared by content length
    total = sum(weights) or 1
    base = 1.5
    extra = max(0.0, CONTENT_W_CM - base * ncols)
    widths = [Cm(base + extra * w / total) for w in weights]
    for idx, w in enumerate(widths):
        table.columns[idx].width = w

    hdr_cells = table.rows[0].cells
    for idx, text in enumerate(header):
        hdr_cells[idx].width = widths[idx]
        p = hdr_cells[idx].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, text, size=9, color=RGBColor(0x0B, 0x4F, 0x57), bold=True)
        cell_shade(hdr_cells[idx], HEADER_FILL)
    repeat_header(table.rows[0])

    for r in body:
        cells = table.add_row().cells
        for idx in range(ncols):
            cells[idx].width = widths[idx]
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, r[idx] if idx < len(r) else "", size=9)

    no_split_rows(table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.35


def add_cover(doc, title, meta_rows):
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("产品需求文档")
    set_run_font(run, CN_FONT, 11, bold=True, color=ACCENT)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    set_run_font(run, CN_FONT, 30, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    rule_runs = p.add_run("　")
    set_run_font(rule_runs, CN_FONT, 6, color=ACCENT)
    para_border(p, "bottom", ACCENT_HEX, size=12, space=1)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    set_run_font(run, CN_FONT, 9, color=MUTED)

    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table_borders(table, "DDE5EA")
    cell_margins(table, 60, 60, 100, 100)
    for k, v in meta_rows:
        cells = table.add_row().cells
        cells[0].width = Cm(4.2)
        cells[1].width = Cm(CONTENT_W_CM - 4.2)
        cell_shade(cells[0], "F7F9FA")
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        add_inline(p0, k, size=9, color=MUTED, bold=True)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        add_inline(p1, v, size=9)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    meta_lookup = dict(meta_rows)
    version = meta_lookup.get("文档版本", "v1.0")
    run = p.add_run(f"古墨　·　2026-10-20　·　{version}")
    set_run_font(run, CN_FONT, 9.5, color=MUTED)

    doc.add_page_break()


def add_header_footer(section):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("SKYRUNNER　·　产品需求文档")
    set_run_font(run, CN_FONT, 8.5, color=MUTED)
    para_border(hp, "bottom", RULE_HEX, size=4, space=4)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("— ")
    set_run_font(r1, CN_FONT, 8.5, color=MUTED)
    add_field(fp, " PAGE ")
    r2 = fp.add_run(" —")
    set_run_font(r2, CN_FONT, 8.5, color=MUTED)


def main():
    md = open(SRC, encoding="utf-8").read()
    # drop YAML front matter (kept in the .md for agent consumption)
    md = re.sub(r"\A---\r?\n.*?\r?\n---\r?\n", "", md, count=1, flags=re.S)
    blocks = parse_blocks(md)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    normal.paragraph_format.line_spacing = 1.55
    normal.paragraph_format.space_after = Pt(6)

    title = TITLE_ARG or "SKYRUNNER"
    meta_rows = []
    start = 0
    if blocks and blocks[0][0] == "h1":
        if not TITLE_ARG:
            title = blocks[0][1]
        start = 1
    if start < len(blocks) and blocks[start][0] == "table":
        rows = blocks[start][1]
        body = rows[1:]
        for r in body:
            cells = split_table_row(r)
            if is_separator(cells) or len(cells) < 2:
                continue
            meta_rows.append((cells[0], cells[1]))
        start += 1

    sec.different_first_page_header_footer = True

    add_cover(doc, title, meta_rows if meta_rows else [("文档", "PRD")])
    add_header_footer(sec)

    for kind, payload in blocks[start:]:
        if kind == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(payload)
            set_run_font(run, CN_FONT, 16, bold=True, color=ACCENT)
            para_border(p, "bottom", RULE_HEX, size=6, space=4)

        elif kind == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(payload)
            set_run_font(run, CN_FONT, 12, bold=True, color=INK)

        elif kind == "h4":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(payload)
            set_run_font(run, CN_FONT, 10.5, bold=True, color=INK)

        elif kind == "p":
            p = doc.add_paragraph()
            add_inline(p, payload)

        elif kind == "bullet":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run("•　")
            set_run_font(run, CN_FONT, 10.5, color=ACCENT, bold=True)
            add_inline(p, payload)

        elif kind == "num":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.9)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, payload)

        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            para_shade(p, CODE_FILL)
            para_border(p, "left", ACCENT_HEX, size=18, space=6)
            add_inline(p, payload, size=10, color=RGBColor(0x3C, 0x4A, 0x58))

        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.2
            para_shade(p, CODE_FILL)
            run = p.add_run(payload)
            set_run_font(run, MONO_FONT, 8.5, color=RGBColor(0x22, 0x2B, 0x36))

        elif kind == "table":
            build_table(doc, payload)

        elif kind == "hr":
            continue

    doc.core_properties.title = "SKYRUNNER 产品需求文档"
    doc.core_properties.author = "古墨"
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
